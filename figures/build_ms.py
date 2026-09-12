"""
Build the manuscript section as PDF and DOCX, with both figures embedded.

The placeholders in the markdown are replaced by real figure references and
captions. The PDF keeps the figures as vector; the DOCX embeds rasterisations,
because Word's SVG support is uneven across the versions a co-author might open
the file in.
"""
from __future__ import annotations

import html
import re
import subprocess
from pathlib import Path

ROOT = Path(r"C:/Users/USER/Downloads/AIMScribe")
SRC = ROOT / "manuscript" / "software_architecture_section.md"
FIGS = ROOT / "figures"
OUT = ROOT / "manuscript"
CHROME = r"C:/Program Files/Google/Chrome/Application/chrome.exe"

FIG1_CAP = ("**Figure 1 | Architecture block diagram.** Boundaries of the "
            "consulting-room workstation, the third-party electronic health record "
            "(EHR) node and the AIMS Lab backend. The EHR node serves the clinical "
            "page over HTTPS. That page signals the aimscribe.exe tray daemon over a "
            "loopback channel confined to the same machine (Channel A), while "
            "clinical data reaches the backend server to server (Channel B). "
            "Accented pathways are encrypted and mutually authenticated. Audio "
            "travels only from the workstation to the backend, and never towards the "
            "EHR node.")

FIG2_CAP = ("**Figure 2 | Sequence of the acquisition pathway.** Opening a "
            "patient sends two messages at once: a trigger to the local daemon "
            "(1a) and a corroborating notice from the EHR server to the backend "
            "(1b). Acquisition begins immediately, and the backend authorises the "
            "recording only if the daemon's request matches the notice exactly. "
            "Segments are sealed, uploaded and verified by server-side re-hashing "
            "throughout. Local material is deleted as soon as a purge receipt "
            "confirms that a verified copy exists.")


def prepare() -> str:
    md = SRC.read_text(encoding="utf-8")
    md = re.sub(r"\[Insert Figure X:.*?\]",
                f"![fig](ms_fig1_architecture.svg)\n\n{FIG1_CAP}", md, flags=re.S)
    md = re.sub(r"\[Insert Figure Y:.*?\]",
                f"![fig](ms_fig2_sequence.svg)\n\n{FIG2_CAP}", md, flags=re.S)
    md = md.replace("[Insert Figure X", "[UNREPLACED X").replace(
        "[Insert Figure Y", "[UNREPLACED Y")
    return md


def inline(t: str) -> str:
    t = html.escape(t, quote=False)
    t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
    t = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", t)
    return t


def to_html(md: str) -> str:
    out, blocks = [], [b for b in md.split("\n\n") if b.strip()]
    i = 0
    while i < len(blocks):
        b = blocks[i].strip()
        m = re.match(r"^!\[fig\]\(([^)]+)\)$", b)
        if m:
            svg = (FIGS / m.group(1)).read_text(encoding="utf-8").strip()
            cap = ""
            if i + 1 < len(blocks) and blocks[i + 1].lstrip().startswith("**Figure"):
                cap = inline(re.sub(r"\s+", " ", blocks[i + 1]).strip())
                i += 1
            out.append(f"<figure>{svg}<figcaption>{cap}</figcaption></figure>")
        elif b.startswith("### "):
            out.append(f"<h3>{inline(b[4:].strip())}</h3>")
        elif b.startswith("## "):
            out.append(f"<h2>{inline(b[3:].strip())}</h2>")
        else:
            out.append("<p>" + inline(re.sub(r"\s+", " ", b).strip()) + "</p>")
        i += 1
    return "\n".join(out)


CSS = """
@page { size: A4; margin: 22mm 20mm; }
body { font: 11.5pt/1.75 "Times New Roman",Georgia,serif; color:#111;
       margin:0; text-align:justify; }
h2 { font-family:"Helvetica Neue",Helvetica,Arial,sans-serif; font-size:15pt;
     margin:0 0 6mm; text-align:left; }
h3 { font-family:"Helvetica Neue",Helvetica,Arial,sans-serif; font-size:11.5pt;
     margin:7mm 0 2.5mm; text-align:left; page-break-after:avoid; }
p { margin:0 0 3mm; orphans:3; widows:3; }
code { font-family:Consolas,monospace; font-size:.86em; }
figure { margin:7mm 0; page-break-inside:avoid; text-align:center; }
figure svg { width:100%; height:auto; max-height:150mm; }
figcaption { font-size:9.5pt; line-height:1.45; text-align:justify;
             margin-top:3mm; color:#1c1c1c; }
@media screen { body { max-width:190mm; margin:0 auto; padding:16mm 10mm; } }
"""


def main() -> None:
    md = prepare()
    doc = ("<!doctype html><html lang='en'><head><meta charset='utf-8'>"
           "<title>Software Architecture and Data Collection System</title>"
           f"<style>{CSS}</style></head><body>{to_html(md)}</body></html>")
    hp = OUT / "AIMScribe_Software_Architecture_Section.html"
    hp.write_text(doc, encoding="utf-8")

    pdf = OUT / "AIMScribe_Software_Architecture_Section.pdf"
    subprocess.run([CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={pdf}", hp.as_uri()],
                   capture_output=True, timeout=120)
    print(f"  PDF   {pdf.name}  {pdf.stat().st_size/1024:.0f} KB")

    # --- rasterise for Word
    pngs = {}
    for svg in ("ms_fig1_architecture.svg", "ms_fig2_sequence.svg"):
        # Take the size from the viewBox: a hardcoded window leaves the figure
        # letterboxed inside a wider canvas, which then shrinks in Word.
        vb = re.search(r'viewBox="0 0 (\d+) (\d+)"',
                       (FIGS / svg).read_text(encoding="utf-8"))
        w, h = int(vb.group(1)), int(vb.group(2))
        png = OUT / (Path(svg).stem + ".png")
        subprocess.run([CHROME, "--headless", "--disable-gpu",
                        f"--screenshot={png}", f"--window-size={w},{h}",
                        "--default-background-color=FFFFFFFF",
                        "--force-device-scale-factor=2.5",
                        (FIGS / svg).as_uri()], capture_output=True, timeout=120)
        pngs[svg] = png
        print(f"  PNG   {png.name}  {png.stat().st_size/1024:.0f} KB")

    # --- DOCX
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(6)

    def rich(p, text):
        """Render **bold**, `code` and plain runs into a paragraph."""
        for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not tok:
                continue
            if tok.startswith("**"):
                p.add_run(tok[2:-2]).bold = True
            elif tok.startswith("`"):
                r = p.add_run(tok[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt(10.5)
            else:
                p.add_run(tok)

    blocks = [b for b in md.split("\n\n") if b.strip()]
    i = 0
    while i < len(blocks):
        b = blocks[i].strip()
        m = re.match(r"^!\[fig\]\(([^)]+)\)$", b)
        if m:
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(pngs[m.group(1)]), width=Inches(5.4))
            if i + 1 < len(blocks) and blocks[i + 1].lstrip().startswith("**Figure"):
                cap = d.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                cap.paragraph_format.line_spacing = 1.0
                rich(cap, re.sub(r"\s+", " ", blocks[i + 1]).strip())
                for r in cap.runs:
                    r.font.size = Pt(10)
                    r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                i += 1
        elif b.startswith("## "):
            d.add_heading(b[3:].strip(), level=1)
        elif b.startswith("### "):
            d.add_heading(b[4:].strip(), level=2)
        else:
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            rich(p, re.sub(r"\s+", " ", b).strip())
        i += 1

    dp = OUT / "AIMScribe_Software_Architecture_Section.docx"
    d.save(str(dp))
    print(f"  DOCX  {dp.name}  {dp.stat().st_size/1024:.0f} KB")


if __name__ == "__main__":
    main()
