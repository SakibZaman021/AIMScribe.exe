"""Build the CMED integration guide as PDF and DOCX. Local only - not published."""
from __future__ import annotations

import html
import re
import subprocess
from pathlib import Path

ROOT = Path(r"C:/Users/USER/Downloads/AIMScribe")
SRC = ROOT / "cmed" / "CMED_INTEGRATION_GUIDE.md"
FIGS = ROOT / "figures"
OUT = ROOT / "cmed"
CHROME = r"C:/Program Files/Google/Chrome/Application/chrome.exe"


def inline(t: str) -> str:
    t = html.escape(t, quote=False)
    t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
    t = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", t)
    t = re.sub(r"(?<![\w*])\*([^*\n]+)\*(?![\w*])", r"<em>\1</em>", t)
    return t


def convert(md: str, svg_inline=True) -> str:
    lines = md.split("\n")
    out, i, n = [], 0, len(lines)
    while i < n:
        ln = lines[i]
        if ln.startswith("```"):
            j, body = i + 1, []
            while j < n and not lines[j].startswith("```"):
                body.append(lines[j]); j += 1
            out.append("<pre><code>" + html.escape("\n".join(body)) + "</code></pre>")
            i = j + 1; continue
        m = re.match(r"^!\[fig\]\(([^)]+)\)\s*$", ln)
        if m:
            cap, k = "", i + 1
            while k < n and not lines[k].strip():
                k += 1
            if k < n and lines[k].startswith("**Figure"):
                buf = []
                while k < n and lines[k].strip():
                    buf.append(lines[k]); k += 1
                cap = inline(re.sub(r"\s+", " ", " ".join(buf))); i = k - 1
            svg = (FIGS / m.group(1)).read_text(encoding="utf-8").strip()
            out.append(f"<figure>{svg}<figcaption>{cap}</figcaption></figure>")
            i += 1; continue
        m = re.match(r"(#{1,6})\s+(.*)$", ln)
        if m:
            lv = len(m.group(1))
            out.append(f"<h{lv}>{inline(m.group(2))}</h{lv}>"); i += 1; continue
        if re.match(r"^-{3,}\s*$", ln):
            out.append("<hr>"); i += 1; continue
        if ln.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|\s*$", lines[i + 1]):
            head = [c.strip() for c in ln.strip().strip("|").split("|")]
            rows, j = [], i + 2
            while j < n and lines[j].startswith("|"):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")]); j += 1
            t = ["<div class='tw'><table><thead><tr>"]
            t += [f"<th>{inline(c)}</th>" for c in head]
            t.append("</tr></thead><tbody>")
            for r in rows:
                t.append("<tr>" + "".join(f"<td>{inline(c)}</td>" for c in r) + "</tr>")
            t.append("</tbody></table></div>")
            out.append("".join(t)); i = j; continue
        if re.match(r"^\s*[-*]\s+", ln) or re.match(r"^\s*\d+\.\s+", ln):
            ordered = bool(re.match(r"^\s*\d+\.\s+", ln))
            items, j = [], i
            while j < n and (re.match(r"^\s*[-*]\s+", lines[j]) or
                             re.match(r"^\s*\d+\.\s+", lines[j])):
                items.append(re.sub(r"^\s*(?:[-*]|\d+\.)\s+", "", lines[j])); j += 1
            tag = "ol" if ordered else "ul"
            out.append(f"<{tag}>" + "".join(f"<li>{inline(x)}</li>" for x in items) + f"</{tag}>")
            i = j; continue
        if not ln.strip():
            i += 1; continue
        para, j = [], i
        while j < n and lines[j].strip() and not re.match(
                r"^(#{1,6}\s|\||>|```|!\[|-{3,}\s*$|\s*[-*]\s+|\s*\d+\.\s+)", lines[j]):
            para.append(lines[j]); j += 1
        out.append("<p>" + inline(" ".join(para)) + "</p>")
        i = j if j > i else i + 1
    return "\n".join(out)


CSS = """
@page { size: A4; margin: 18mm 17mm; }
* { box-sizing: border-box; }
body { font: 10.8pt/1.6 Georgia,"Times New Roman",serif; color:#141414; margin:0; }
h1,h2,h3,th { font-family:"Helvetica Neue",Helvetica,Arial,sans-serif; }
h1 { font-size:22pt; margin:0 0 3mm; line-height:1.15; }
h2 { font-size:14pt; margin:9mm 0 3mm; padding-bottom:1.5mm;
     border-bottom:1pt solid #141414; page-break-after:avoid; }
h2:first-of-type { margin-top:5mm; }
h3 { font-size:11pt; margin:5mm 0 2mm; page-break-after:avoid; }
p { margin:0 0 2.8mm; }
hr { border:0; border-top:.5pt solid #ccc; margin:5mm 0; }
code { font-family:Consolas,monospace; font-size:.85em; background:#f2f0ef;
       padding:.3mm 1mm; border-radius:1mm; }
pre { background:#f7f5f4; border:.5pt solid #ddd8d6; border-radius:1mm;
      padding:2.5mm 3mm; page-break-inside:avoid; overflow-x:auto; }
pre code { background:none; padding:0; font-size:8.2pt; line-height:1.4; }
.tw { overflow-x:auto; margin:0 0 3.5mm; page-break-inside:avoid; }
table { border-collapse:collapse; width:100%; font-size:9pt; }
th,td { border:.5pt solid #cfcac8; padding:1.5mm 2mm; text-align:left;
        vertical-align:top; }
th { background:#f2efee; font-weight:600; font-size:8.6pt; }
ul,ol { margin:0 0 3mm; padding-left:6mm; }
li { margin-bottom:1.2mm; }
figure { margin:6mm 0; page-break-inside:avoid; text-align:center; }
figure svg { width:100%; height:auto; max-height:200mm; }
figcaption { font-size:9pt; text-align:left; margin-top:2.5mm; line-height:1.45; }
@media screen { body { max-width:196mm; margin:0 auto; padding:14mm 10mm; } }
"""


def main() -> None:
    md = SRC.read_text(encoding="utf-8")
    doc = ("<!doctype html><html lang='en'><head><meta charset='utf-8'>"
           "<title>AIMScribe — Integration Guide for CMED</title>"
           f"<style>{CSS}</style></head><body>{convert(md)}</body></html>")
    hp = OUT / "CMED_Integration_Guide.html"
    hp.write_text(doc, encoding="utf-8")

    pdf = OUT / "CMED_Integration_Guide.pdf"
    subprocess.run([CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={pdf}", hp.as_uri()], capture_output=True, timeout=120)
    print(f"  PDF   {pdf.name}  {pdf.stat().st_size/1024:.0f} KB")

    pngs = {}
    for svg in re.findall(r"!\[fig\]\(([^)]+)\)", md):
        vb = re.search(r'viewBox="0 0 (\d+) (\d+)"', (FIGS / svg).read_text(encoding="utf-8"))
        w, h = int(vb.group(1)), int(vb.group(2))
        png = OUT / (Path(svg).stem + ".png")
        subprocess.run([CHROME, "--headless", "--disable-gpu", f"--screenshot={png}",
                        f"--window-size={w},{h}", "--default-background-color=FFFFFFFF",
                        "--force-device-scale-factor=2.5", (FIGS / svg).as_uri()],
                       capture_output=True, timeout=120)
        pngs[svg] = png

    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15

    def rich(p, text, size=None):
        for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not tok:
                continue
            if tok.startswith("**"):
                r = p.add_run(tok[2:-2]); r.bold = True
            elif tok.startswith("`"):
                r = p.add_run(tok[1:-1]); r.font.name = "Consolas"
                r.font.size = Pt((size or 11) - 1)
            else:
                r = p.add_run(tok)
            if size:
                r.font.size = Pt(size)

    lines = md.split("\n"); i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            j, body = i + 1, []
            while j < len(lines) and not lines[j].startswith("```"):
                body.append(lines[j]); j += 1
            p = d.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run("\n".join(body)); r.font.name = "Consolas"; r.font.size = Pt(8.5)
            i = j + 1; continue
        m = re.match(r"^!\[fig\]\(([^)]+)\)\s*$", ln)
        if m:
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(pngs[m.group(1)]), width=Inches(5.6))
            k = i + 1
            while k < len(lines) and not lines[k].strip():
                k += 1
            if k < len(lines) and lines[k].startswith("**Figure"):
                buf = []
                while k < len(lines) and lines[k].strip():
                    buf.append(lines[k].strip()); k += 1
                cap = d.add_paragraph(); rich(cap, " ".join(buf), size=9); i = k - 1
            i += 1; continue
        m = re.match(r"(#{1,3})\s+(.*)$", ln)
        if m:
            d.add_heading(re.sub(r"[*`]", "", m.group(2)), level=len(m.group(1)))
            i += 1; continue
        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|\s*$", lines[i + 1]):
            head = [c.strip() for c in ln.strip().strip("|").split("|")]
            rows, j = [], i + 2
            while j < len(lines) and lines[j].startswith("|"):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")]); j += 1
            t = d.add_table(rows=1, cols=len(head)); t.style = "Table Grid"
            for c, h in zip(t.rows[0].cells, head):
                c.text = ""; rich(c.paragraphs[0], h, size=9)
                for r in c.paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                cells = t.add_row().cells
                for c, v in zip(cells, row):
                    c.text = ""; rich(c.paragraphs[0], v, size=9)
            d.add_paragraph(); i = j; continue
        if re.match(r"^\s*[-*]\s+", ln) or re.match(r"^\s*\d+\.\s+", ln):
            ordered = bool(re.match(r"^\s*\d+\.\s+", ln))
            while i < len(lines) and (re.match(r"^\s*[-*]\s+", lines[i]) or
                                      re.match(r"^\s*\d+\.\s+", lines[i])):
                p = d.add_paragraph(style="List Number" if ordered else "List Bullet")
                rich(p, re.sub(r"^\s*(?:[-*]|\d+\.)\s+", "", lines[i])); i += 1
            continue
        if re.match(r"^-{3,}\s*$", ln) or not ln.strip():
            i += 1; continue
        para, j = [], i
        while j < len(lines) and lines[j].strip() and not re.match(
                r"^(#{1,6}\s|\||```|!\[|-{3,}\s*$|\s*[-*]\s+|\s*\d+\.\s+)", lines[j]):
            para.append(lines[j]); j += 1
        rich(d.add_paragraph(), " ".join(para)); i = j if j > i else i + 1

    dp = OUT / "CMED_Integration_Guide.docx"
    d.save(str(dp))
    print(f"  DOCX  {dp.name}  {dp.stat().st_size/1024:.0f} KB")

    body = re.sub(r"^#+ .*$", "", md, flags=re.M)
    print(f"  words {len(re.sub(r'`{3}.*?`{3}', '', body, flags=re.S).split())}")


if __name__ == "__main__":
    main()
